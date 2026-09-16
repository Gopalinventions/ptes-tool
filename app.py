import io
import folium
import geopandas as gpd
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from designer_integration import design_geometry, designer_html
from energy_balance import DEFAULT_DEMAND_SHARES, EnergySystemInputs, simulate_monthly_balance
from hourly_dispatch import run_hourly_dispatch
from thermal_ui import render_thermal
from folium.features import GeoJsonPopup
from folium.plugins import Draw
from shapely.affinity import rotate, translate
from shapely.geometry import LineString, Point, box
from streamlit_folium import st_folium
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from calculations import (assess_main_pipe, darcy_weisbach_pressure_loss,
                          engineering_suitability_score,
                          geodetic_pressure_correction, required_flow,
                          size_connection_pipe, size_storage_from_demand,
                          storage_capacity, suitability_score)
from gis_analysis import clean_geometry, local_metric_crs, nearby_demand, nearest_pipe_connection, split_layers
from spatial_analysis import (containing_parcel_measurements,
                              intersecting_feature_count,
                              nearest_distance_m,
                              nearest_point_attribute,
                              overlap_area_m2, read_uploaded_vector)
from weather_analysis import (TEMPERATURE_HINTS, TIMESTAMP_HINTS,
                              build_weather_demand_profile, monthly_summary,
                              read_weather_table, seasonal_summary,
                              suggested_column)

st.set_page_config(page_title="PTES Comparison Tool", page_icon="♨️", layout="wide")
st.title("PTES Location and Network Integration Tool")
st.caption("Place and compare Storage A, B and C. Click buildings and pipes to inspect nPro data.")

if "candidates" not in st.session_state:
    st.session_state.candidates = {}
if "energy_hub_index" not in st.session_state:
    st.session_state.energy_hub_index = None

BUILDING_FIELDS = {
    "b_building_name": "Building", "b_addr_street": "Street",
    "b_addr_house_number": "House number", "b_building_type": "Building type",
    "b_year_constr": "Construction year", "b_floor_area_m2": "Floor area [m²]",
    "b_heat_import_sum_MWh": "Heat import [MWh/year]",
    "b_space_heat_sum_MWh": "Space heat [MWh/year]",
    "b_heat_import_max_kW": "Peak heat [kW]",
    "b_supply_temp_heat_degC": "Supply temperature [°C]",
}
PIPE_FIELDS = {
    "p_pipe_id": "Pipe ID", "p_pipe_model": "Pipe model", "p_diameter_DN": "DN",
    "p_length_m": "Length [m]", "p_is_existing_option": "Existing pipe",
    "p_mean_temp_supply_degC": "Supply temperature [°C]",
    "p_mean_temp_return_degC": "Return temperature [°C]",
    "p_abs_pres_supply_bar": "Supply pressure [bar]",
    "p_abs_pres_return_bar": "Return pressure [bar]",
    "p_abs_pres_supply_sim_bar": "Simulated supply pressure [bar]",
    "p_abs_pres_return_sim_bar": "Simulated return pressure [bar]",
    "p_pressure_loss_rel_sim_Pa_m": "Simulated loss [Pa/m]",
    "p_max_flow_velocity_m_s": "Maximum velocity [m/s]",
    "p_max_power_possible_kW": "Possible power [kW]",
}
COLORS = {"Storage A": "green", "Storage B": "orange", "Storage C": "purple"}
SCORING_CRITERIA = [
    "Energy-hub suitability", "Demand and storage performance", "Network connection",
    "Hydraulic compatibility", "Land and construction fit",
    "GIS/environmental constraints", "Data confidence",
]
THEMES = {
    "Standard network": None,
    "Supply pressure [bar]": "p_abs_pres_supply_sim_bar",
    "Return pressure [bar]": "p_abs_pres_return_sim_bar",
    "Supply temperature [°C]": "p_mean_temp_supply_degC",
    "Return temperature [°C]": "p_mean_temp_return_degC",
    "Pressure loss [Pa/m]": "p_pressure_loss_rel_sim_Pa_m",
    "Flow velocity [m/s]": "p_max_flow_velocity_m_s",
    "Power reserve [kW]": "p_power_reserve_abs_kW",
}


@st.cache_data(show_spinner="Reading energy-profile file…")
def read_energy_csv(file_bytes: bytes) -> pd.DataFrame:
    """Read comma- or semicolon-separated planning data without assuming a vendor."""
    return pd.read_csv(io.BytesIO(file_bytes), sep=None, engine="python")


@st.cache_data(show_spinner="Reading Excel energy-profile workbook…")
def excel_sheet_names(file_bytes: bytes) -> list[str]:
    return pd.ExcelFile(io.BytesIO(file_bytes)).sheet_names


@st.cache_data(show_spinner="Reading Excel energy-profile sheet…")
def read_energy_excel(file_bytes: bytes, sheet_name: str, header_row: int) -> pd.DataFrame:
    return pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet_name, header=header_row - 1)


def numeric_series(frame: pd.DataFrame, column: str) -> pd.Series:
    """Accept decimal commas as well as decimal points in uploaded CSV files."""
    return pd.to_numeric(frame[column].astype(str).str.replace(",", ".", regex=False), errors="coerce")


def hourly_energy_profile(timestamps: pd.Series, energy: pd.Series, name: str) -> pd.DataFrame:
    """Normalise uploaded interval-energy data onto unique UTC-naive hourly timestamps."""
    stamp = pd.to_datetime(timestamps, errors="coerce", utc=True).dt.tz_convert(None).dt.floor("h")
    table = pd.DataFrame({"Timestamp": stamp, name: energy}).dropna()
    return table.groupby("Timestamp", as_index=False)[name].sum()


def default_hourly_demand(index: pd.DatetimeIndex, annual_mwh: float) -> pd.Series:
    """Fallback monthly demand distribution; never presented as an uploaded profile."""
    hours_by_month = pd.Series(index.month).value_counts().to_dict()
    return pd.Series([annual_mwh * DEFAULT_DEMAND_SHARES[stamp.month - 1] / hours_by_month[stamp.month]
                      for stamp in index], index=index)


def seasonal_storage_cycle(frame: pd.DataFrame, start_month: int) -> pd.DataFrame:
    """Show one representative charge/discharge cycle from start month to next year."""
    result = frame.copy()
    result["Timestamp"] = pd.to_datetime(result["Timestamp"])
    result.loc[result["Timestamp"].dt.month < start_month, "Timestamp"] = (
        result.loc[result["Timestamp"].dt.month < start_month, "Timestamp"] + pd.DateOffset(years=1)
    )
    return result.sort_values("Timestamp").reset_index(drop=True)


def json_safe(layer):
    layer = layer.copy()
    for col in layer.columns:
        if col != layer.geometry.name:
            layer[col] = layer[col].map(lambda x: None if pd.isna(x) else x if isinstance(x, (str, int, float, bool)) else str(x))
    return layer


@st.cache_data(show_spinner="Reading and projecting GIS data…")
def read_vector_bytes(file_bytes, target_crs=None, assume_wgs84=False):
    """Cache expensive GeoJSON parsing across Streamlit widget reruns."""
    layer = gpd.read_file(io.BytesIO(file_bytes))
    if layer.crs is None:
        if not assume_wgs84:
            raise ValueError("The file has no CRS metadata.")
        layer = layer.set_crs(4326)
    layer = clean_geometry(layer)
    return layer if target_crs is None else layer.to_crs(target_crs)


def nearby_map_layers(layers, candidates, metric_crs, radius_m):
    """Limit heavy contextual layers to the active candidate investigation area."""
    if not candidates:
        return {name: None for name in layers}
    points = gpd.GeoSeries(
        [Point(item["Longitude"], item["Latitude"]) for item in candidates], crs=4326
    ).to_crs(metric_crs)
    investigation_area = points.union_all().buffer(radius_m)
    nearby = {}
    for name, layer in layers.items():
        nearby[name] = None if layer is None else layer[layer.geometry.intersects(investigation_area)].copy()
    return nearby


def excel_workbook(candidate_results, data_register, seasonal_results=None, monthly_results=None):
    """Create one workbook with all important outputs on one summary sheet."""
    def excel_safe_frame(frame):
        """Return an Excel-compatible copy; Excel cannot store timezone-aware datetimes."""
        safe = frame.copy()
        for column in safe.columns:
            if isinstance(safe[column].dtype, pd.DatetimeTZDtype):
                safe[column] = safe[column].dt.tz_convert("UTC").dt.tz_localize(None)
            elif safe[column].dtype == "object":
                safe[column] = safe[column].map(
                    lambda value: value.tz_convert("UTC").tz_localize(None)
                    if isinstance(value, pd.Timestamp) and value.tzinfo is not None else value
                )
        return safe

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        sheet_name = "Engineering Summary"
        row = 1
        candidate_results = candidate_results.copy()
        candidate_results["Candidate"] = pd.Categorical(
            candidate_results["Candidate"], ["Storage A", "Storage B", "Storage C"], ordered=True
        )
        candidate_results = candidate_results.sort_values("Candidate").reset_index(drop=True)
        candidate_results["Candidate"] = candidate_results["Candidate"].astype(str)
        excel_safe_frame(candidate_results).to_excel(writer, sheet_name=sheet_name, index=False, startrow=row)
        sheet = writer.sheets[sheet_name]
        sheet.cell(row=1, column=1, value="PTES candidate comparison")
        row += len(candidate_results) + 4
        sheet.cell(row=row, column=1, value="GIS data register")
        excel_safe_frame(data_register).to_excel(writer, sheet_name=sheet_name, index=False, startrow=row)
        row += len(data_register) + 3
        if seasonal_results is not None and not seasonal_results.empty:
            sheet.cell(row=row, column=1, value="Weather-derived seasonal demand")
            excel_safe_frame(seasonal_results).to_excel(writer, sheet_name=sheet_name, index=False, startrow=row)
            row += len(seasonal_results) + 3
        if monthly_results is not None and not monthly_results.empty:
            sheet.cell(row=row, column=1, value="Weather-derived monthly demand")
            excel_safe_frame(monthly_results).to_excel(writer, sheet_name=sheet_name, index=False, startrow=row)
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "B3"
            sheet.auto_filter.ref = f"A2:{sheet.cell(row=2, column=len(candidate_results.columns)).coordinate}"
            sheet.sheet_view.showGridLines = False
            sheet.row_dimensions[1].height = 25
            sheet.row_dimensions[2].height = 42
            for cell in sheet[1]:
                cell.font = cell.font.copy(bold=True, size=14)
            for cell in sheet[2]:
                cell.font = cell.font.copy(bold=True, color="FFFFFF")
                cell.fill = cell.fill.copy(fill_type="solid", fgColor="1F4E78")
                cell.alignment = cell.alignment.copy(wrap_text=True, vertical="center")
            for column in sheet.columns:
                width = min(45, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
                sheet.column_dimensions[column[0].column_letter].width = width
    return output.getvalue()


def project_excel_workbook(inputs, energy_monthly, hourly_result, geometry, candidate_results,
                           data_register, seasonal_results=None, monthly_weather=None):
    """Create a readable multi-tab PTES planning workbook from the active tool results."""
    def safe(frame):
        copy = frame.copy()
        for column in copy.columns:
            if isinstance(copy[column].dtype, pd.DatetimeTZDtype):
                copy[column] = copy[column].dt.tz_convert("UTC").dt.tz_localize(None)
        return copy

    def write_table(writer, sheet_name, title, explanation, frame):
        sheet = writer.book.create_sheet(sheet_name)
        sheet["A1"] = title
        sheet["A2"] = explanation
        safe(frame).to_excel(writer, sheet_name=sheet_name, index=False, startrow=3)
        sheet.freeze_panes = "A5"
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[2].height = 35
        for cell in sheet[1]:
            cell.font = Font(bold=True, size=15, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(frame.columns)))
        sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        for cell in sheet[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.auto_filter.ref = f"A4:{sheet.cell(row=4 + len(frame), column=max(1, len(frame.columns))).coordinate}"
        for column in sheet.columns:
            letter = column[0].column_letter
            width = min(34, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[letter].width = width
        return sheet

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        overview = pd.DataFrame([
            {"Step": "1. Inputs", "Purpose": "Active PTES, solar, BHKW and operating-month assumptions."},
            {"Step": "2. Monthly planning", "Purpose": "Summer price-selected BHKW charging plan and PTES monthly balance."},
            {"Step": "3. Hourly dispatch", "Purpose": "Solar-to-PTES, demand-led winter BHKW, PTES discharge and boiler balance."},
            {"Step": "4. Geometry", "Purpose": "Preliminary dimensions and civil quantities for the selected PTES."},
            {"Step": "5. GIS screening", "Purpose": "Candidate, network and constraint screening results."},
            {"Step": "6. Weather", "Purpose": "Uploaded or derived weather-demand summaries, when available."},
        ])
        write_table(writer, "Read me", "PTES planning workbook", "This workbook records the active tool settings and results. It is a planning and screening export, not a final hydraulic, geotechnical or detailed-design calculation.", overview)
        write_table(writer, "Inputs", "Active inputs", "Values used when this workbook was downloaded. Change values in the Streamlit tool and download a new workbook to compare cases.", inputs)
        monthly_sheet = write_table(writer, "Monthly balance", "Step 2 — monthly PTES planning", "Solar is allocated to PTES first. BHKW values in this tab are the summer, price-selected PTES-charging plan only; winter BHKW operation is shown in Hourly dispatch.", energy_monthly)
        if not energy_monthly.empty:
            chart = BarChart()
            chart.title = "Monthly PTES charge and discharge"
            chart.y_axis.title = "MWh"
            chart.x_axis.title = "Month"
            names = list(energy_monthly.columns)
            selected = [names.index(name) + 1 for name in ["PTES charge [MWh]", "PTES discharge [MWh]"] if name in names]
            if selected:
                for column in selected:
                    chart.add_data(Reference(monthly_sheet, min_col=column, min_row=4, max_row=4 + len(energy_monthly)), titles_from_data=True)
                chart.set_categories(Reference(monthly_sheet, min_col=1, min_row=5, max_row=4 + len(energy_monthly)))
                monthly_sheet.add_chart(chart, "A20")
        if hourly_result is not None and not hourly_result.empty:
            hourly_summary = hourly_result.set_index("Timestamp").resample("ME").sum(numeric_only=True).reset_index()
            hourly_sheet = write_table(writer, "Hourly summary", "Step 3 — monthly summary of hourly dispatch", "This is the operational result. Summer BHKW-to-PTES hours use the price threshold. Winter BHKW direct-network supply follows heat demand and does not use price.", hourly_summary)
            chart = LineChart()
            chart.title = "PTES end-of-month state of charge"
            chart.y_axis.title = "MWh"
            names = list(hourly_summary.columns)
            if "State of charge [MWh]" in names:
                chart.add_data(Reference(hourly_sheet, min_col=names.index("State of charge [MWh]") + 1, min_row=4, max_row=4 + len(hourly_summary)), titles_from_data=True)
                chart.set_categories(Reference(hourly_sheet, min_col=1, min_row=5, max_row=4 + len(hourly_summary)))
                hourly_sheet.add_chart(chart, "A20")
            write_table(writer, "Hourly detail", "Step 3 — hourly dispatch detail", "Each row is one modelled hour. Use this tab for checking BHKW 1, BHKW 2, solar charging, PTES state of charge, discharge and boiler heat.", hourly_result)
        geometry_frame = pd.DataFrame([geometry])
        write_table(writer, "Geometry", "Step 4 — PTES geometry", "Preliminary geometry and quantities. Confirm slopes, liner, cover, groundwater, drainage and civil works with specialist design inputs.", geometry_frame)
        write_table(writer, "GIS screening", "Step 5 — candidate and GIS screening", "Candidate rankings and distances are screening outputs. Missing or approximate GIS layers require verification with authority and survey data.", candidate_results)
        write_table(writer, "GIS register", "Step 5 — GIS data register", "Data availability and source-layer register used in the screening.", data_register)
        if seasonal_results is not None and not seasonal_results.empty:
            write_table(writer, "Weather seasonal", "Step 6 — seasonal weather result", "Seasonal weather-derived heat-demand summary.", seasonal_results)
        if monthly_weather is not None and not monthly_weather.empty:
            write_table(writer, "Weather monthly", "Step 6 — monthly weather result", "Monthly weather-derived heat-demand summary.", monthly_weather)
    return output.getvalue()


def hourly_dispatch_excel_workbook(hourly_result, monthly_energy, monthly_soc, inputs):
    """Export the operating model as an understandable workbook, not a raw CSV."""
    output = io.BytesIO()
    readme = pd.DataFrame([
        {"Section": "Summer BHKW charging", "Explanation": "Only May–September BHKW-to-PTES operation uses the selected day-ahead price threshold."},
        {"Section": "Winter BHKW heat", "Explanation": "October–April BHKW direct-network heat follows demand. It is not selected by electricity price."},
        {"Section": "Solar thermal", "Explanation": "Solar thermal charges PTES first. It is curtailed only when PTES is full."},
        {"Section": "PTES discharge", "Explanation": "PTES supplies heat only in selected discharge months and only while stored heat is available."},
        {"Section": "Boiler heat", "Explanation": "Boiler heat is the remaining network demand after direct sources and PTES discharge."},
    ])
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="Read me", index=False)
        inputs.to_excel(writer, sheet_name="Active inputs", index=False)
        monthly_energy.reset_index().to_excel(writer, sheet_name="Monthly operations", index=False)
        monthly_soc.reset_index().to_excel(writer, sheet_name="PTES state of charge", index=False)
        hourly_result.to_excel(writer, sheet_name="Hourly dispatch", index=False)
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.sheet_view.showGridLines = False
            sheet.row_dimensions[1].height = 32
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(32, max(13, max(len(str(cell.value or "")) for cell in column) + 2))
        operations = writer.sheets["Monthly operations"]
        columns = list(monthly_energy.reset_index().columns)
        chart_columns = [name for name in ["Solar to PTES [MWh]", "BHKW 1 to PTES [MWh]", "BHKW 2 to PTES [MWh]", "BHKW 1 direct network [MWh]", "BHKW 2 direct network [MWh]", "PTES discharge [MWh]", "Boiler heat [MWh]"] if name in columns]
        if chart_columns:
            chart = BarChart()
            chart.type = "col"
            chart.style = 10
            chart.title = "Monthly energy flows"
            chart.y_axis.title = "MWh"
            chart.x_axis.title = "Month"
            for name in chart_columns:
                chart.add_data(Reference(operations, min_col=columns.index(name) + 1, min_row=1, max_row=len(monthly_energy) + 1), titles_from_data=True)
            chart.set_categories(Reference(operations, min_col=1, min_row=2, max_row=len(monthly_energy) + 1))
            operations.add_chart(chart, "A18")
        soc_sheet = writer.sheets["PTES state of charge"]
        soc_columns = list(monthly_soc.reset_index().columns)
        if "State of charge [MWh]" in soc_columns:
            chart = LineChart()
            chart.title = "PTES end-of-month state of charge"
            chart.y_axis.title = "MWh"
            chart.add_data(Reference(soc_sheet, min_col=soc_columns.index("State of charge [MWh]") + 1, min_row=1, max_row=len(monthly_soc) + 1), titles_from_data=True)
            chart.set_categories(Reference(soc_sheet, min_col=1, min_row=2, max_row=len(monthly_soc) + 1))
            soc_sheet.add_chart(chart, "A18")
    return output.getvalue()


def add_layer(fmap, layer, name, definitions, style, highlight):
    fields = [f for f in definitions if f in layer.columns]
    popup = GeoJsonPopup(fields, [definitions[f] for f in fields], localize=True) if fields else None
    folium.GeoJson(layer, name=name, style_function=lambda _: style,
                   highlight_function=lambda _: highlight, popup=popup).add_to(fmap)


def add_network(fmap, buildings, pipes, theme_name="Standard network",
                building_demand_field=None, demand_classes=12):
    """Draw network data and a 10–12 class annual building-demand legend."""
    demand_palette = [
        "#ffffcc", "#ffeda0", "#fed976", "#feb24c", "#fd8d3c", "#fc4e2a",
        "#ef3b2c", "#e31a1c", "#bd0026", "#99000d", "#67000d", "#3f0010",
    ]
    if building_demand_field and building_demand_field in buildings.columns:
        values = pd.to_numeric(buildings[building_demand_field], errors="coerce")
        low, high = float(values.min()), float(values.max())
        palette = [
            demand_palette[round(i * (len(demand_palette) - 1) / max(demand_classes - 1, 1))]
            for i in range(demand_classes)
        ]

        def building_style(feature):
            try:
                value = float(feature["properties"].get(building_demand_field))
                fraction = 0.0 if high == low else (value - low) / (high - low)
                index = min(demand_classes - 1, max(0, int(fraction * demand_classes)))
                color = palette[index]
            except (TypeError, ValueError):
                color = "#B8B8B8"
            return {"color": "#555", "weight": 1, "fillColor": color, "fillOpacity": .72}

        building_fields = dict(BUILDING_FIELDS)
        building_fields.setdefault(building_demand_field, "Annual heat demand [MWh/year]")
        popup_fields = [field for field in building_fields if field in buildings.columns]
        folium.GeoJson(
            buildings, name="Buildings — annual heat demand", style_function=building_style,
            highlight_function=lambda _: {"color": "#00FFFF", "weight": 3, "fillOpacity": .9},
            popup=GeoJsonPopup(popup_fields, [building_fields[field] for field in popup_fields], localize=True),
        ).add_to(fmap)
        intervals = []
        for index, color in enumerate(palette):
            start = low + (high - low) * index / demand_classes
            end = low + (high - low) * (index + 1) / demand_classes
            intervals.append(
                f"<div><span style='display:inline-block;width:13px;height:10px;background:{color};"
                f"margin-right:5px'></span>{start:,.0f}–{end:,.0f}</div>"
            )
        building_legend = (
            "<div style='position:fixed;bottom:24px;left:24px;z-index:9998;background:white;"
            "padding:10px 12px;border:1px solid #777;border-radius:4px;font-size:11px;"
            "max-height:310px;overflow:auto'><b>Building annual heat demand</b><br>"
            "MWh/year<br>" + "".join(intervals) +
            "<div><span style='display:inline-block;width:13px;height:10px;background:#B8B8B8;"
            "margin-right:5px'></span>No value</div></div>"
        )
        fmap.get_root().html.add_child(folium.Element(building_legend))
    else:
        add_layer(fmap, buildings, "Buildings — click for details", BUILDING_FIELDS,
                  {"color": "#555", "weight": 1, "fillColor": "#F3A712", "fillOpacity": .28},
                  {"weight": 3, "fillOpacity": .5})
    field = THEMES.get(theme_name)
    if field and field in pipes.columns:
        values = pd.to_numeric(pipes[field], errors="coerce")
        low, high = values.min(), values.max()
        colors = ["#313695", "#74add1", "#ffffbf", "#f46d43", "#a50026"]
        def themed_style(feature):
            try:
                value = float(feature["properties"].get(field))
                fraction = 0 if high == low else (value - low) / (high - low)
                color = colors[min(4, max(0, int(fraction * 4.999)))]
            except (TypeError, ValueError):
                color = "#999999"
            return {"color": color, "weight": 5, "opacity": .9}
        fields = [f for f in PIPE_FIELDS if f in pipes.columns]
        folium.GeoJson(pipes, name=theme_name, style_function=themed_style,
                       highlight_function=lambda _: {"color": "#00FFFF", "weight": 8},
                       popup=GeoJsonPopup(fields, [PIPE_FIELDS[f] for f in fields], localize=True)).add_to(fmap)
        legend = f"<div style='position:fixed;top:18px;left:55px;z-index:9999;background:white;padding:10px;border:1px solid #777'><b>{theme_name}</b><br>Low&nbsp; <span style='color:#313695'>■</span> <span style='color:#74add1'>■</span> <span style='color:#ffffbf'>■</span> <span style='color:#f46d43'>■</span> <span style='color:#a50026'>■</span>&nbsp; High<br>{low:.2f} – {high:.2f}</div>"
        fmap.get_root().html.add_child(folium.Element(legend))
    else:
        add_layer(fmap, pipes, "Pipes — click for details", PIPE_FIELDS,
                  {"color": "#1677FF", "weight": 4}, {"color": "#00B8D9", "weight": 7})


def add_basemaps(fmap):
    """Add street and satellite backgrounds to Streamlit and exported HTML maps."""
    folium.TileLayer("OpenStreetMap", name="Street map", control=True).add_to(fmap)
    folium.TileLayer(
        tiles="https://server.arcgisonline.com/ArcGIS/rest/services/World_Imagery/MapServer/tile/{z}/{y}/{x}",
        attr="Esri, Maxar, Earthstar Geographics, and the GIS User Community",
        name="Satellite imagery",
        overlay=False,
        control=True,
    ).add_to(fmap)


def add_optional_layers(fmap, layers):
    styles = {
        "Candidate parcels": {"color": "#6B4F2A", "weight": 2, "fillOpacity": .05},
        "Groundwater": {"color": "#00A6D6", "weight": 3, "fillColor": "#00A6D6", "fillOpacity": .25},
        "Flood zones": {"color": "#0066CC", "weight": 2, "fillColor": "#4DA6FF", "fillOpacity": .30},
        "Protected areas": {"color": "#167A3E", "weight": 2, "fillColor": "#4CAF50", "fillOpacity": .25},
        "Roads": {"color": "#555555", "weight": 3},
        "Known utilities": {"color": "#D81B60", "weight": 4, "dashArray": "6 4"},
    }
    for name, layer in layers.items():
        if layer is None or layer.empty:
            continue
        wgs = json_safe(layer.to_crs(4326))
        fields = [c for c in wgs.columns if c != "geometry"][:12]
        popup = GeoJsonPopup(fields, fields, localize=True) if fields else None
        folium.GeoJson(wgs, name=name, style_function=lambda _, s=styles[name]: s,
                       highlight_function=lambda _: {"weight": 6}, popup=popup).add_to(fmap)


def marker(fmap, name, lat, lon, result=None):
    text = f"<b>{name}</b><br>Latitude: {lat:.6f}<br>Longitude: {lon:.6f}"
    if result:
        branch_dn = result.get("Branch DN", result.get("DN", "Not available"))
        main_dn = result.get("Main DN", "Not available")
        text += (
            f"<br>Score: {result['Score [%]']:.1f}%"
            f"<br>Screening class: {result.get('Screening classification', 'Not assessed')}"
            f"<br>Connection: {result['Connection [m]']:.1f} m"
            f"<br>Main pipe DN: {main_dn}"
            f"<br>PTES branch DN: {branch_dn}"
        )
    folium.Marker([lat, lon], tooltip=name, popup=folium.Popup(text, max_width=320),
                  icon=folium.Icon(color=COLORS[name], icon="info-sign")).add_to(fmap)


def energy_hub_marker(fmap, lat, lon, label):
    folium.Marker(
        [lat, lon], tooltip=f"Energy hub: {label}",
        popup=folium.Popup(f"<b>Selected energy hub</b><br>{label}", max_width=320),
        icon=folium.Icon(color="red", icon="wrench", prefix="fa"),
    ).add_to(fmap)


with st.sidebar:
    with st.expander("1 · nPro network data", expanded=True):
        uploaded = st.file_uploader("Upload nPro GeoJSON", type=["geojson", "json"])
        st.caption("Buildings and existing district-heating pipes are read from this file.")
        hub_selector_placeholder = st.empty()
    with st.expander("2 · Storage size and geometry", expanded=True):
        st.caption("Sets the water volume and physical pit dimensions for drawings, GIS and thermal layers.")
        storage_volume_mode = st.selectbox("Storage-volume method", ["Use available volume", "Calculate from demand"])
        available_storage_volume = st.number_input(
            "Target water volume [m³]", 1.0, value=125000.0,
            disabled=storage_volume_mode != "Use available volume",
        )
        depth = st.number_input("Total pit depth, including freeboard [m]", 1.0, value=15.0)
    with st.expander("2b · Network demand and static capacity estimate", expanded=storage_volume_mode == "Calculate from demand"):
        st.caption("Annual demand supplies GIS/weather context. Coverage sizes the pit only in demand-based mode. These are static estimates, not hourly operating conditions.")
        annual_demand = st.number_input("Annual system heat demand [MWh/year]", 1.0, value=20000.0)
        storage_type = st.selectbox("Storage type", ["Seasonal", "Weekly", "Daily"])
        coverage = st.slider("Annual demand allocated to storage [%]", 1.0, 100.0, 30.0,
                             disabled=storage_volume_mode != "Calculate from demand")
        tmax = st.number_input("Static capacity: hot reference [°C]", value=90.0)
        tmin = st.number_input("Static capacity: cold reference [°C]", value=15.0)
        efficiency = st.slider("Static usable-capacity factor (not simulated losses)", .5, 1.0, .8)
        st.caption("This factor is used only for static sizing/capacity. The thermal simulation calculates boundary losses separately and does not multiply by this factor.")
        reference_demand = st.number_input("Reference demand [MWh/year]", 1.0, value=10000.0)
    with st.expander("2c · Interlinked energy-system planning", expanded=True):
        st.caption("Monthly screen: solar charges PTES first and BHKW values show only the summer price-selected PTES-charging plan. Use Step 0b for the separate demand-led winter BHKW network operation.")
        solar_thermal_kw = st.number_input("Solar thermal nominal capacity [kWth]", 0.0, value=3300.0)
        solar_specific_yield = st.number_input(
            "Solar net specific yield [kWhth/kWth/year]", 0.0, value=1030.0,
            help="Enter a site-specific annual net yield from your solar model. The default is only a WÜST planning-screen value.",
        )
        solar_monthly_profile = None
        solar_hourly_profile = None
        solar_profile_file = st.file_uploader(
            "Optional solar heat-energy profile (CSV or Excel)", type=["csv", "txt", "xlsx", "xls"], key="solar_profile"
        )
        if solar_profile_file is not None:
            try:
                solar_bytes = solar_profile_file.getvalue()
                if solar_profile_file.name.lower().endswith((".xlsx", ".xls")):
                    sheet_names = excel_sheet_names(solar_bytes)
                    default_sheet = sheet_names.index("Hourly") if "Hourly" in sheet_names else 0
                    solar_sheet = st.selectbox("Solar workbook sheet", sheet_names, index=default_sheet, key="solar_sheet")
                    default_header = 5 if solar_sheet == "Hourly" else 1
                    solar_header_row = st.number_input("Header row in solar sheet", min_value=1, value=default_header, key="solar_header_row")
                    solar_frame = read_energy_excel(solar_bytes, solar_sheet, int(solar_header_row))
                else:
                    solar_frame = read_energy_csv(solar_bytes)
                solar_columns = list(solar_frame.columns)
                time_default = next((i for i, col in enumerate(solar_columns) if any(word in col.lower() for word in ("time", "date", "stamp"))), 0)
                energy_default = next(
                    (i for i, col in enumerate(solar_columns) if "net balance" in col.lower()),
                    next((i for i, col in enumerate(solar_columns) if any(word in col.lower() for word in ("heat", "energy", "solar", "mwh", "kwh"))), min(1, len(solar_columns) - 1)),
                )
                solar_time_col = st.selectbox("Solar profile timestamp column", solar_columns, index=time_default, key="solar_time_col")
                solar_energy_col = st.selectbox("Solar interval-energy column", solar_columns, index=energy_default, key="solar_energy_col")
                unit_default = 1 if "kwh" in solar_energy_col.lower() else 0
                solar_unit = st.selectbox("Solar interval-energy unit", ["MWh", "kWh"], index=unit_default, key="solar_energy_unit")
                solar_time = pd.to_datetime(solar_frame[solar_time_col], errors="coerce")
                solar_values = numeric_series(solar_frame, solar_energy_col)
                multiplier = 0.001 if solar_unit == "kWh" else 1.0
                solar_hourly_profile = hourly_energy_profile(solar_time, solar_values * multiplier, "Solar [MWh]")
                grouped = pd.DataFrame({"month": solar_time.dt.month, "energy": solar_values * multiplier}).dropna().groupby("month")["energy"].sum()
                solar_monthly_profile = tuple(float(grouped.get(month, 0.0)) for month in range(1, 13))
                st.success(f"Solar profile accepted: {sum(solar_monthly_profile):,.0f} MWh net heat across 12 months.")
            except (ValueError, KeyError, pd.errors.ParserError) as exc:
                st.warning(f"Solar profile could not be read: {exc}. The annual-yield method remains active.")
        st.markdown("**BHKW design — direct energy-hub supply and optional PTES charging**")
        bhkw1_electrical_kw = st.number_input("BHKW 1 electrical capacity [kWe]", 0.0, value=975.0)
        bhkw1_thermal_kw = st.number_input("BHKW 1 thermal capacity [kWth]", 0.0, value=1195.0)
        bhkw2_electrical_kw = st.number_input("BHKW 2 electrical capacity [kWe]", 0.0, value=975.0)
        bhkw2_thermal_kw = st.number_input("BHKW 2 thermal capacity [kWth]", 0.0, value=1195.0)
        bhkw_electrical_kw = bhkw1_electrical_kw + bhkw2_electrical_kw
        bhkw_thermal_kw = bhkw1_thermal_kw + bhkw2_thermal_kw
        st.caption(f"Combined BHKW capacity used in the simple dispatch: {bhkw_electrical_kw:,.0f} kWe and {bhkw_thermal_kw:,.0f} kWth.")
        bhkw_fuel_per_mwh_e = st.number_input("BHKW fuel input per electricity output [MWhfuel/MWhe]", 0.1, value=2.4728)
        bhkw_month_labels = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
        bhkw_direct_labels = st.multiselect(
            "BHKW direct-network supply months (heat-demand led; no price rule)",
            bhkw_month_labels,
            default=["Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr"],
        )
        bhkw_charge_labels = st.multiselect(
            "BHKW-to-PTES charging months (summer price rule applies)",
            bhkw_month_labels,
            default=["May", "Jun", "Jul", "Aug", "Sep"],
        )
        ptes_discharge_labels = st.multiselect("PTES discharge months", bhkw_month_labels,
                                               default=["Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr"])
        storage_cycle_view = st.selectbox(
            "PTES cycle view", ["May–April seasonal storage year", "January–December calendar year"],
            help="May–April is recommended: solar/BHKW summer charging happens before winter discharge through the following April.",
        )
        bhkw_direct_months = tuple(bhkw_month_labels.index(month) + 1 for month in bhkw_direct_labels)
        bhkw_charge_months = tuple(bhkw_month_labels.index(month) + 1 for month in bhkw_charge_labels)
        ptes_discharge_months = tuple(bhkw_month_labels.index(month) + 1 for month in ptes_discharge_labels)
        bhkw_active_months = tuple(sorted(set(bhkw_direct_months) | set(bhkw_charge_months)))
        bhkw_mode = st.radio("BHKW operating-hours method", ["Manual planning hours", "2025 day-ahead price threshold"], horizontal=True)
        day_ahead_summary = None
        day_ahead_hourly = None
        price_threshold = 100.0
        if bhkw_mode == "Manual planning hours":
            bhkw_summer_hours = st.number_input("BHKW selected summer operating hours [h/year]", 0.0, value=0.0)
        else:
            day_ahead_file = st.file_uploader("Upload day-ahead electricity-price CSV", type=["csv", "txt"], key="day_ahead_prices")
            price_threshold = st.number_input("BHKW minimum day-ahead price [€/MWh]", value=100.0)
            bhkw_summer_hours = 0.0
            if day_ahead_file is not None:
                try:
                    price_frame = read_energy_csv(day_ahead_file.getvalue())
                    price_columns = list(price_frame.columns)
                    price_time_default = next((i for i, col in enumerate(price_columns) if any(word in col.lower() for word in ("start", "time", "date"))), 0)
                    price_value_default = next((i for i, col in enumerate(price_columns) if any(word in col.lower() for word in ("price", "germany", "luxembourg", "eur"))), min(1, len(price_columns) - 1))
                    price_time_col = st.selectbox("Day-ahead timestamp column", price_columns, index=price_time_default, key="price_time_col")
                    price_value_col = st.selectbox("Day-ahead price column", price_columns, index=price_value_default, key="price_value_col")
                    price_time = pd.to_datetime(price_frame[price_time_col], errors="coerce")
                    prices = numeric_series(price_frame, price_value_col)
                    # Day-ahead prices decide only optional BHKW-to-PTES summer operation.
                    # Direct network heat is dispatched later from the heat demand, not the price.
                    selected = (price_time.dt.month.isin(bhkw_charge_months) & (prices >= price_threshold)).fillna(False)
                    bhkw_summer_hours = float(selected.sum())
                    day_ahead_summary = {
                        "hours": bhkw_summer_hours,
                        "mean_price": float(prices[selected].mean()) if selected.any() else 0.0,
                        "revenue_eur": float((prices[selected] * bhkw_electrical_kw / 1000).sum()),
                    }
                    day_ahead_hourly = hourly_energy_profile(price_time, prices, "Price [€/MWh]")
                    st.success(f"Selected summer BHKW charging hours: {bhkw_summer_hours:,.0f} h; mean price: {day_ahead_summary['mean_price']:,.1f} €/MWh.")
                except (ValueError, KeyError, pd.errors.ParserError) as exc:
                    st.warning(f"Day-ahead price file could not be read: {exc}. Upload the CSV again or use manual hours.")
            else:
                st.info("Upload the day-ahead CSV to calculate BHKW operating hours from the selected price threshold.")
        show_advanced_energy = st.checkbox("Show advanced heat-pump, waste-heat and storage settings", value=False)
        heat_pump_thermal_kw = heat_pump_summer_hours = waste_heat_kw = waste_heat_summer_hours = 0.0
        heat_pump_cop, monthly_storage_loss = 3.0, 0.0
        hourly_initial_soc, hourly_hp_max_price = 0.0, 80.0
        if show_advanced_energy:
            heat_pump_thermal_kw = st.number_input("Heat-pump thermal capacity combined [kWth]", 0.0, value=0.0)
            heat_pump_summer_hours = st.number_input("Heat-pump summer operating hours [h/year]", 0.0, value=0.0)
            heat_pump_cop = st.number_input("Heat-pump seasonal COP", min_value=0.1, value=3.0)
            waste_heat_kw = st.number_input("Waste-heat available capacity [kWth]", 0.0, value=0.0)
            waste_heat_summer_hours = st.number_input("Waste-heat summer availability [h/year]", 0.0, value=0.0)
            monthly_storage_loss = st.number_input("PTES monthly standing loss [%]", min_value=0.0, max_value=99.0, value=0.0)
            hourly_initial_soc = st.slider("Initial PTES state of charge at first hour [%]", 0.0, 100.0, 0.0) / 100
            hourly_hp_max_price = st.number_input("Heat-pump maximum electricity price [€/MWh]", value=80.0)
            st.caption("For a normal full-year simulation keep initial PTES state of charge at 0%. The heat pump runs only below its selected electricity-price limit.")
        demand_hourly_profile = None
        demand_profile_file = st.file_uploader("Optional hourly heat-demand profile CSV", type=["csv", "txt"], key="demand_profile")
        if demand_profile_file is not None:
            try:
                demand_frame = read_energy_csv(demand_profile_file.getvalue())
                demand_columns = list(demand_frame.columns)
                demand_time_default = next((i for i, col in enumerate(demand_columns) if any(word in col.lower() for word in ("time", "date", "stamp"))), 0)
                demand_energy_default = next((i for i, col in enumerate(demand_columns) if any(word in col.lower() for word in ("heat", "demand", "energy", "mwh", "kwh"))), min(1, len(demand_columns) - 1))
                demand_time_col = st.selectbox("Demand-profile timestamp column", demand_columns, index=demand_time_default, key="demand_time_col")
                demand_energy_col = st.selectbox("Demand interval-energy column", demand_columns, index=demand_energy_default, key="demand_energy_col")
                demand_unit = st.selectbox("Demand interval-energy unit", ["MWh", "kWh"], key="demand_energy_unit")
                demand_multiplier = 0.001 if demand_unit == "kWh" else 1.0
                demand_hourly_profile = hourly_energy_profile(
                    demand_frame[demand_time_col], numeric_series(demand_frame, demand_energy_col) * demand_multiplier,
                    "Demand [MWh]",
                )
                st.success(f"Demand profile accepted: {demand_hourly_profile['Demand [MWh]'].sum():,.0f} MWh across {len(demand_hourly_profile):,} hours.")
            except (ValueError, KeyError, pd.errors.ParserError) as exc:
                st.warning(f"Demand profile could not be read: {exc}. A monthly-distribution fallback will be labelled in the output.")
    with st.expander("3 · Connection-pipe design point"):
        st.caption("One hydraulic operating point for DN and pressure-loss screening. This is NOT a charging schedule. Run hourly operation in Step 3 of the main page.")
        operating_mode = st.selectbox("Hydraulic screening mode", ["Charging", "Discharging", "Idle"])
        power = st.number_input("Pipe-design heat transfer [kW]", 1.0, value=3300.0)
        delta_t = st.number_input("Pipe supply–return design ΔT [K]", 1.0, value=30.0)
        branch_dn_options = [50, 65, 80, 100, 125, 150, 200, 250, 300]
        minimum_branch_dn = st.selectbox("Minimum PTES branch DN", branch_dn_options,
                                         index=branch_dn_options.index(150))
        maximum_branch_dn = st.selectbox("Maximum PTES branch DN", branch_dn_options,
                                         index=branch_dn_options.index(250))
    with st.expander("4 · Official GIS layers"):
        st.caption("Upload authority-supplied GeoJSON layers. Missing layers are reported as not assessed.")
        parcels_file = st.file_uploader("Candidate parcels", type=["geojson", "json"], key="parcels")
        groundwater_file = st.file_uploader("Groundwater observations", type=["geojson", "json"], key="groundwater")
        flood_file = st.file_uploader("Flood zones", type=["geojson", "json"], key="flood")
        protected_file = st.file_uploader("Protected areas", type=["geojson", "json"], key="protected")
        roads_file = st.file_uploader("Roads", type=["geojson", "json"], key="roads")
        utilities_file = st.file_uploader("Known utilities", type=["geojson", "json"], key="utilities")
    with st.expander("5 · Weather-derived demand"):
        weather_source = st.selectbox("Weather-data source", ["Not supplied", "DWD TRY", "ERA5-Land", "Other hourly CSV"])
        weather_file = st.file_uploader("Weather CSV or TXT", type=["csv", "txt"], key="weather")
        dhw_share = st.number_input("Domestic-hot-water share [%]", 0.0, 99.0, value=12.0)
        heating_limit = st.number_input("Heating-limit temperature [°C]", value=15.0)
    with st.expander("6 · Advanced geometry and pressure"):
        freeboard = st.number_input("Freeboard below rim [m]", 0.0, value=2.5)
        side_slope = st.number_input("Side slope H:V", 0.0, value=1.5)
        aspect_ratio = st.number_input("Bottom length-to-width ratio", .2, value=1.16)
        st.caption("Illustrative design values. Water depth = total depth − freeboard. Slopes require geotechnical verification.")
        rotation = st.number_input("Footprint rotation [degrees]", value=0.0)
        elevation_offset = st.number_input("Candidate elevation above nearest pipe [m]", value=0.0)
        embankment_width = st.number_input(
            "Permanent perimeter / embankment allowance [m]", 0.0, value=5.0,
            help=("Initial planning allowance outside the excavated rim for the cover edge, "
                  "anchor trench, drainage detail and embankment crest. It is not a slope-stability design."),
        )
        construction_clearance = st.number_input(
            "Temporary construction working clearance [m]", 0.0, value=8.0,
            help=("Initial planning corridor outside the permanent perimeter for access, liner installation "
                  "and construction plant. Confirm it with the construction method and site logistics plan."),
        )
        st.caption(
            "Compact preliminary layout: 5 m permanent perimeter + 8 m temporary working corridor. "
            "The former 15 m + 20 m screening buffers were intentionally conservative. Increase these "
            "values where geotechnics, haul roads, cranes, drainage or temporary soil stockpiles require it."
        )
        show_layout_options = st.checkbox("Show preliminary pump / diffuser layout options", value=False)
        pump_side, pump_chamber_length, pump_chamber_width = "East", 12.0, 8.0
        drainage_well_count, diffuser_velocity_limit = 2, 0.03
        if show_layout_options:
            st.caption("These values draw a concept layout only. They are not pump, diffuser, drainage or civil-work specifications.")
            pump_side = st.selectbox("Pump chamber side", ["East", "West", "North", "South"], index=0)
            pump_chamber_length = st.number_input("Pump chamber length [m]", 1.0, value=12.0)
            pump_chamber_width = st.number_input("Pump chamber width [m]", 1.0, value=8.0)
            drainage_well_count = st.number_input("Preliminary drainage / monitoring points", 0, 8, value=2, step=1)
            diffuser_velocity_limit = st.number_input(
                "Diffuser outlet-velocity screening limit [m/s]", 0.005, 0.10, value=0.03, step=0.005,
                help="Low velocity is used as a screening principle to limit mixing. The final diffuser geometry requires hydraulic design.",
            )
        parcel_radius = st.number_input("Nearby GIS investigation radius [m]", 100.0, value=1000.0,
                                        help="Only nearby parcel and context features are drawn on the map.")
    with st.expander("7 · Suitability percentage"):
        selected_scoring_criteria = st.multiselect(
            "Criteria included in percentage",
            SCORING_CRITERIA,
            default=SCORING_CRITERIA,
            help=("Choose 1–7 criteria. The established engineering weights of the selected "
                  "criteria are automatically re-normalised to 100%."),
        )
        st.caption(
            "Unselected criteria are excluded, not scored as zero. Selected base weights are "
            "re-normalised so their applied weights total 100%."
        )
    with st.expander("8 · Optional material quantities"):
        st.caption("Editable examples only. Enter supplier specifications before using quantities for construction LCA.")
        liner_thickness = st.number_input("Liner thickness [mm]", .1, value=2.0)
        liner_density = st.number_input("Liner density [kg/m³]", 1.0, value=940.0)
        liner_allowance = st.number_input("Liner overlaps / waste [%]", 0.0, value=5.0)
        cover_thickness = st.number_input("Cover insulation thickness [mm]", .1, value=240.0)
        cover_density = st.number_input("Cover insulation density [kg/m³]", 1.0, value=30.0)

st.sidebar.markdown(
    """
    <div style="margin-top:1.5rem;padding:0.85rem 0.9rem;border:1px solid #d7dee8;
                border-left:5px solid #d97706;border-radius:0.55rem;background:#f8fafc;
                color:#172033;line-height:1.35;">
      <div style="font-size:1.05rem;font-weight:700;">♨️ PTES Energy Screening</div>
      <div style="font-size:0.78rem;margin-top:0.45rem;">Developed by<br>
        <strong>Kuruba Pujari Gopal</strong>
      </div>
      <div style="font-size:0.76rem;margin-top:0.4rem;color:#52606d;">
        Rother und Partner Ingenieurgesellschaft
      </div>
    </div>
    """,
    unsafe_allow_html=True,
)

try:
    if storage_volume_mode == "Calculate from demand":
        storage = size_storage_from_demand(annual_demand, coverage, storage_type, tmax, tmin, efficiency)
    else:
        capacity = storage_capacity(available_storage_volume, tmax, tmin, efficiency)
        cycles = {"Seasonal": 1, "Weekly": 52, "Daily": 365}[storage_type]
        storage = dict(volume_m3=available_storage_volume,
                       energy_per_cycle_mwh=capacity["useful_capacity_mwh"],
                       annual_shifted_mwh=capacity["useful_capacity_mwh"]*cycles,
                       cycles_per_year=cycles, delta_t_k=capacity["delta_t_k"])
    design_flow = required_flow(power, delta_t)
    layout = {
        "pumpSide": pump_side,
        "pumpLength": pump_chamber_length,
        "pumpWidth": pump_chamber_width,
        "drainageWells": int(drainage_well_count),
        "designFlowM3h": design_flow["volume_flow_m3_h"],
        "diffuserVelocityMps": diffuser_velocity_limit,
        "diffuserAreaM2": design_flow["volume_flow_m3_s"] / diffuser_velocity_limit,
        "minimumDN": minimum_branch_dn,
        "maximumDN": maximum_branch_dn,
    }
    design_inputs, design_model, geometry = design_geometry(
        storage["volume_m3"], depth, side_slope, freeboard, aspect_ratio,
        linerThickness=liner_thickness, linerDensity=liner_density,
        allowance=liner_allowance, coverThickness=cover_thickness, coverDensity=cover_density,
        permanentPerimeter=embankment_width, temporaryWorking=construction_clearance,
        layout=layout)
    design_document = designer_html(design_inputs, design_model)
    energy_inputs = EnergySystemInputs(
        annual_demand_mwh=annual_demand,
        storage_volume_m3=storage["volume_m3"], hot_c=tmax, cold_c=tmin,
        usable_capacity_factor=efficiency,
        solar_net_annual_mwh=solar_thermal_kw * solar_specific_yield / 1000,
        solar_monthly_mwh=solar_monthly_profile,
        bhkw_electrical_kw=bhkw_electrical_kw,
        bhkw_thermal_kw=bhkw_thermal_kw, bhkw_summer_hours=bhkw_summer_hours,
        bhkw_operating_months=bhkw_charge_months,
        heat_pump_thermal_kw=heat_pump_thermal_kw,
        heat_pump_summer_hours=heat_pump_summer_hours, heat_pump_cop=heat_pump_cop,
        waste_heat_kw=waste_heat_kw, waste_heat_summer_hours=waste_heat_summer_hours,
        monthly_storage_loss_percent=monthly_storage_loss,
    )
    energy_balance = simulate_monthly_balance(energy_inputs)
except (ValueError, OSError) as exc:
    st.error(f"Design could not be prepared: {exc}")
    st.stop()

st.subheader("Step 0 · Interlinked PTES energy-system planning")
st.caption("This generic planning screen is independent of one project site. It links the selected PTES volume with solar, BHKW, heat-pump, waste-heat and demand inputs. Monthly default profiles are clearly preliminary until you upload hourly profiles.")
energy_col1, energy_col2, energy_col3, energy_col4 = st.columns(4)
energy_col1.metric("Selected PTES capacity", f"{energy_balance['capacity_mwh']:,.0f} MWh")
energy_col2.metric("Maximum PTES state of charge", f"{energy_balance['maximum_soc_mwh']:,.0f} MWh")
energy_col3.metric("PTES winter discharge", f"{energy_balance['totals']['PTES discharge [MWh]']:,.0f} MWh")
energy_col4.metric("Remaining boiler heat", f"{energy_balance['totals']['Remaining boiler heat [MWh]']:,.0f} MWh")
if day_ahead_summary is not None:
    market_col1, market_col2, market_col3 = st.columns(3)
    market_col1.metric("BHKW price-selected hours", f"{day_ahead_summary['hours']:,.0f} h")
    market_col2.metric("Selected mean price", f"{day_ahead_summary['mean_price']:,.1f} €/MWh")
    market_col3.metric("Gross BHKW electricity revenue", f"€{day_ahead_summary['revenue_eur']:,.0f}")
    st.caption("Gross market revenue only. Gas, O&M, starts, demand acceptance and PTES limits are assessed separately.")
with st.expander("Open linked charging, discharging and source balance", expanded=True):
    energy_frame = pd.DataFrame(energy_balance["rows"])
    bhkw1_share = 0.0 if bhkw_thermal_kw == 0 else bhkw1_thermal_kw / bhkw_thermal_kw
    energy_frame["BHKW 1 heat [MWh]"] = energy_frame["BHKW heat [MWh]"] * bhkw1_share
    energy_frame["BHKW 2 heat [MWh]"] = energy_frame["BHKW heat [MWh]"] - energy_frame["BHKW 1 heat [MWh]"]
    chart_frame = energy_frame.set_index("Month")
    chart_frame.index = pd.date_range("2025-01-01", periods=12, freq="MS")
    chart_frame.index.name = "Month"
    monthly_left, monthly_right = st.columns(2)
    with monthly_left:
        st.caption("Monthly sources and PTES charge/discharge")
        st.bar_chart(chart_frame[["Solar to PTES [MWh]", "BHKW 1 heat [MWh]", "BHKW 2 heat [MWh]", "PTES charge [MWh]", "PTES discharge [MWh]", "Remaining boiler heat [MWh]"]], use_container_width=True)
    with monthly_right:
        st.caption("Monthly PTES state of charge")
        st.line_chart(chart_frame[["PTES state of charge [MWh]"]], use_container_width=True)
    st.dataframe(energy_frame.round(1), hide_index=True, use_container_width=True)
    st.download_button(
        "Download energy-balance CSV", energy_frame.to_csv(index=False).encode("utf-8"),
        "ptes_interlinked_energy_balance.csv", "text/csv",
    )
    st.info(energy_balance["model_status"])

st.subheader("Step 0b · Hourly dispatch — uploaded profiles")
st.caption("Solar and demand are aligned to the uploaded day-ahead hourly timestamps. Solar thermal charges PTES first and is curtailed only when PTES is full; it does not directly supply the network in this operating concept.")
hourly_result = None
if day_ahead_hourly is None:
    st.info("For hourly BHKW dispatch, select ‘2025 day-ahead price threshold’ in Section 2c and upload the day-ahead price CSV.")
else:
    hourly_input = day_ahead_hourly.copy()
    if solar_hourly_profile is None:
        hourly_input["Solar [MWh]"] = 0.0
        solar_status = "No hourly solar file uploaded: solar is set to zero in the hourly engine."
    else:
        hourly_input = hourly_input.merge(solar_hourly_profile, on="Timestamp", how="left")
        hourly_input["Solar [MWh]"] = hourly_input["Solar [MWh]"].fillna(0.0)
        solar_status = "Uploaded solar hourly profile used."
    if demand_hourly_profile is None:
        hourly_input["Demand [MWh]"] = default_hourly_demand(pd.DatetimeIndex(hourly_input["Timestamp"]), annual_demand).to_numpy()
        demand_status = "Fallback monthly demand distribution used — upload hourly demand for a dispatch decision."
    else:
        hourly_input = hourly_input.merge(demand_hourly_profile, on="Timestamp", how="left")
        hourly_input["Demand [MWh]"] = hourly_input["Demand [MWh]"].fillna(0.0)
        demand_status = "Uploaded hourly demand profile used."
    if storage_cycle_view == "May–April seasonal storage year":
        hourly_input = seasonal_storage_cycle(hourly_input, start_month=5)
        cycle_status = "Seasonal cycle shown as May 2025–April 2026; January–April source profiles reuse the representative uploaded-year data."
    else:
        cycle_status = "Calendar-year cycle shown from January–December."
    try:
        hourly_result = run_hourly_dispatch(
            hourly_input[["Timestamp", "Demand [MWh]", "Solar [MWh]", "Price [€/MWh]"]],
            storage_capacity_mwh=energy_balance["capacity_mwh"], initial_soc_fraction=hourly_initial_soc,
            bhkw1_thermal_kw=bhkw1_thermal_kw, bhkw1_electrical_kw=bhkw1_electrical_kw,
            bhkw2_thermal_kw=bhkw2_thermal_kw, bhkw2_electrical_kw=bhkw2_electrical_kw,
            bhkw_price_threshold=price_threshold, bhkw_fuel_per_mwh_e=bhkw_fuel_per_mwh_e,
            bhkw_direct_months=bhkw_direct_months, bhkw_charge_months=bhkw_charge_months,
            ptes_discharge_months=ptes_discharge_months,
            heat_pump_thermal_kw=heat_pump_thermal_kw, heat_pump_cop=heat_pump_cop,
            heat_pump_max_price=hourly_hp_max_price, waste_heat_kw=waste_heat_kw,
            monthly_loss_percent=monthly_storage_loss,
        )
        hourly_metrics = hourly_result.sum(numeric_only=True)
        h1, h2, h3, h4 = st.columns(4)
        h1.metric("PTES charged", f"{hourly_metrics['PTES charge [MWh]']:,.0f} MWh")
        h2.metric("PTES discharged", f"{hourly_metrics['PTES discharge [MWh]']:,.0f} MWh")
        h3.metric("BHKW electricity revenue", f"€{hourly_metrics['BHKW electricity revenue [€]']:,.0f}")
        h4.metric("Boiler heat remaining", f"{hourly_metrics['Boiler heat [MWh]']:,.0f} MWh")
        st.caption(f"{solar_status} {demand_status} {cycle_status}")
        hourly_indexed = hourly_result.set_index("Timestamp")
        # Support an older hourly_dispatch.py during GitHub/Streamlit updates.
        # The fallback keeps the app running and is replaced by exact BHKW 1/2
        # flows as soon as the matching updated dispatch file is deployed.
        legacy_dispatch = False
        if "Solar to PTES [MWh]" not in hourly_indexed:
            hourly_indexed["Solar to PTES [MWh]"] = 0.0
            legacy_dispatch = True
        total_bhkw = hourly_indexed.get("BHKW heat [MWh]", pd.Series(0.0, index=hourly_indexed.index))
        bhkw1_share_hourly = (hourly_indexed.get("BHKW 1 heat [MWh]", pd.Series(0.0, index=hourly_indexed.index))
                              .div(total_bhkw.where(total_bhkw > 0, 1.0))).clip(0.0, 1.0)
        direct_bhkw = hourly_indexed.get("BHKW direct network [MWh]", pd.Series(0.0, index=hourly_indexed.index))
        charge_bhkw = hourly_indexed.get("BHKW to PTES [MWh]", pd.Series(0.0, index=hourly_indexed.index))
        for label, total_flow in [("direct network", direct_bhkw), ("to PTES", charge_bhkw)]:
            one = f"BHKW 1 {label} [MWh]"
            two = f"BHKW 2 {label} [MWh]"
            if one not in hourly_indexed or two not in hourly_indexed:
                hourly_indexed[one] = total_flow * bhkw1_share_hourly
                hourly_indexed[two] = total_flow - hourly_indexed[one]
                legacy_dispatch = True
        if legacy_dispatch:
            st.info("Compatibility view: the deployed dispatch file is older than the interface. BHKW 1/2 split is estimated from each unit's hourly heat share until you upload the matching hourly_dispatch.py file.")
        monthly_energy = hourly_indexed[[
            "Solar [MWh]", "Solar to PTES [MWh]",
            "BHKW 1 direct network [MWh]", "BHKW 2 direct network [MWh]",
            "BHKW 1 to PTES [MWh]", "BHKW 2 to PTES [MWh]",
            "PTES charge [MWh]", "PTES discharge [MWh]", "Boiler heat [MWh]",
        ]].resample("ME").sum()
        monthly_soc = hourly_indexed[["State of charge [MWh]"]].resample("ME").last()
        if monthly_energy[["BHKW 1 direct network [MWh]", "BHKW 2 direct network [MWh]"]].to_numpy().sum() == 0:
            st.warning("Winter BHKW direct-network heat is zero in this run. Check that October–April are selected as BHKW direct-network months and that the hourly demand profile contains winter demand.")
        chart_left, chart_right = st.columns(2)
        with chart_left:
            st.caption("Monthly source and PTES heat flows")
            st.bar_chart(monthly_energy[["Solar to PTES [MWh]", "BHKW 1 direct network [MWh]", "BHKW 2 direct network [MWh]", "BHKW 1 to PTES [MWh]", "BHKW 2 to PTES [MWh]", "PTES discharge [MWh]", "Boiler heat [MWh]"]], use_container_width=True)
        with chart_right:
            st.caption("PTES end-of-month state of charge — not summed")
            st.line_chart(monthly_soc, use_container_width=True)
        st.markdown("**Monthly operating table — BHKW flows shown separately**")
        source_table = monthly_energy[["BHKW 1 direct network [MWh]", "BHKW 2 direct network [MWh]", "BHKW 1 to PTES [MWh]", "BHKW 2 to PTES [MWh]", "Boiler heat [MWh]"]].copy()
        source_table.index = source_table.index.strftime("%b %Y")
        source_table.index.name = "Month"
        storage_table = monthly_energy[["Solar to PTES [MWh]", "PTES charge [MWh]", "PTES discharge [MWh]"]].copy()
        storage_table["PTES state of charge [MWh]"] = monthly_soc["State of charge [MWh]"]
        storage_table.index = storage_table.index.strftime("%b %Y")
        storage_table.index.name = "Month"
        table_left, table_right = st.columns(2)
        with table_left:
            st.caption("BHKW and network heat")
            st.dataframe(source_table.round(1), use_container_width=True)
        with table_right:
            st.caption("Solar and PTES")
            st.dataframe(storage_table.round(1), use_container_width=True)
        hourly_export_inputs = pd.DataFrame([
            {"Input": "Summer BHKW charging price threshold [€/MWh]", "Value": price_threshold},
            {"Input": "BHKW direct-network months", "Value": ", ".join(bhkw_direct_labels)},
            {"Input": "BHKW-to-PTES charging months", "Value": ", ".join(bhkw_charge_labels)},
            {"Input": "PTES discharge months", "Value": ", ".join(ptes_discharge_labels)},
            {"Input": "PTES cycle view", "Value": storage_cycle_view},
            {"Input": "Initial PTES state of charge [%]", "Value": hourly_initial_soc * 100},
        ])
        st.download_button("Download hourly dispatch Excel workbook (tables, charts and explanation)",
                           hourly_dispatch_excel_workbook(hourly_result, monthly_energy, monthly_soc, hourly_export_inputs),
                           "ptes_hourly_dispatch_explained.xlsx",
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        st.download_button("Download hourly dispatch CSV", hourly_result.to_csv(index=False).encode("utf-8"),
                           "ptes_hourly_dispatch.csv", "text/csv")
    except ValueError as exc:
        st.warning(f"Hourly dispatch could not be calculated: {exc}")

if uploaded is None:
    st.subheader("Step 1 · nPro network, energy hub and storage candidate")
    st.info("Upload the nPro GeoJSON in sidebar section 1. Then select the energy hub, place Storage A/B/C and run the nPro/GIS screening. Steps 2 and 3 use the selected storage geometry.")
    st.stop()

try:
    data = read_vector_bytes(uploaded.getvalue(), assume_wgs84=True)
    buildings, pipes = split_layers(data)
    metric_crs = local_metric_crs(data)
    buildings_m, pipes_m = clean_geometry(buildings.to_crs(metric_crs)), clean_geometry(pipes.to_crs(metric_crs))
except Exception as exc:
    st.error(f"Could not prepare GeoJSON: {exc}")
    st.stop()

st.subheader("Step 1 · nPro network, energy hub and storage candidate")
st.caption("Select an energy-hub building and place Storage A, B or C. The map evaluates the route from energy hub → storage → existing network. Storage dimensions are calculated separately in Step 2; nPro does not provide final civil dimensions.")

optional_files = {
    "Candidate parcels": parcels_file,
    "Groundwater": groundwater_file,
    "Flood zones": flood_file,
    "Protected areas": protected_file,
    "Roads": roads_file,
    "Known utilities": utilities_file,
}
optional_layers = {}
for layer_name, layer_file in optional_files.items():
    try:
        optional_layers[layer_name] = None if layer_file is None else read_vector_bytes(
            layer_file.getvalue(), target_crs=str(metric_crs)
        )
    except Exception as exc:
        st.error(f"{layer_name} could not be loaded: {exc}")
        optional_layers[layer_name] = None

data_register = [{"Dataset": "Existing network and buildings", "Source file": uploaded.name,
                  "Features": len(data), "Assessment status": "Loaded"}]
for layer_name, layer_file in optional_files.items():
    layer = optional_layers[layer_name]
    data_register.append({"Dataset": layer_name,
                          "Source file": layer_file.name if layer_file else "Not supplied",
                          "Features": 0 if layer is None else len(layer),
                          "Assessment status": "Not assessed" if layer is None else "Loaded"})

register = pd.DataFrame(data_register)
with st.expander("GIS data register", expanded=True):
    st.caption("This register records which measured datasets are available for the current investigation.")
    st.dataframe(register, hide_index=True, use_container_width=True)

weather_seasonal = pd.DataFrame()
weather_monthly = pd.DataFrame()
if weather_source != "Not supplied" and weather_file is not None:
    try:
        weather_table = read_weather_table(weather_file.getvalue())
        weather_columns = list(weather_table.columns)
        timestamp_guess = suggested_column(weather_columns, TIMESTAMP_HINTS)
        temperature_guess = suggested_column(weather_columns, TEMPERATURE_HINTS)
        st.subheader("Weather-derived seasonal demand")
        wc1, wc2 = st.columns(2)
        timestamp_column = wc1.selectbox(
            "Weather timestamp column", weather_columns,
            index=weather_columns.index(timestamp_guess),
        )
        temperature_column = wc2.selectbox(
            "Outdoor-temperature column", weather_columns,
            index=weather_columns.index(temperature_guess),
        )
        weather_profile = build_weather_demand_profile(
            weather_table, timestamp_column, temperature_column,
            annual_demand, dhw_share, heating_limit,
        )
        if storage_volume_mode == "Use available volume":
            comparison_volume = available_storage_volume
        else:
            comparison_volume = size_storage_from_demand(
                annual_demand, coverage, storage_type, tmax, tmin, efficiency
            )["volume_m3"]
        comparison_capacity = storage_capacity(comparison_volume, tmax, tmin, efficiency)
        weather_seasonal = seasonal_summary(weather_profile, comparison_capacity["useful_capacity_mwh"])
        weather_monthly = monthly_summary(weather_profile)
        w1, w2, w3 = st.columns(3)
        w1.metric("Profile records", f"{len(weather_profile):,}")
        w2.metric("Compared PTES volume", f"{comparison_volume:,.0f} m³")
        w3.metric("Usable stored heat", f"{comparison_capacity['useful_capacity_mwh']:,.0f} MWh")
        st.dataframe(weather_seasonal.round(2), hide_index=True, use_container_width=True)
        chart_data = weather_monthly.set_index("month")[["demand_mwh"]]
        st.bar_chart(chart_data, y_label="Heat demand [MWh/month]")
        st.caption(
            f"{weather_source} temperature-derived synthetic demand profile. "
            "It is normalized to the entered annual demand and is not a measured network load profile."
        )
    except Exception as exc:
        st.error(f"Weather profile could not be calculated: {exc}")

groundwater_layer = optional_layers["Groundwater"]
groundwater_fields = [] if groundwater_layer is None else [c for c in groundwater_layer.columns if c != "geometry"]
groundwater_level_field = st.selectbox(
    "Groundwater-level/elevation field",
    ["Not supplied"] + groundwater_fields,
    disabled=not groundwater_fields,
)

demand_options = [c for c in buildings_m.columns if c != "geometry"]
preferred = next((c for c in ("b_heat_import_sum_MWh", "b_space_heat_sum_MWh") if c in demand_options), demand_options[0])
demand_col = st.selectbox("Annual heat-demand column", demand_options, index=demand_options.index(preferred))
building_demand_classes = st.select_slider(
    "Building heat-demand legend classes", options=[10, 11, 12], value=12,
    help="The selected number of colour classes is included in the downloadable HTML map.",
)
buildings_m[demand_col] = pd.to_numeric(buildings_m[demand_col], errors="coerce").fillna(0)
buildings_m = buildings_m.reset_index(drop=True)
hub_name_field = next((c for c in ("b_building_name", "b_addr_street", "b_building_type")
                       if c in buildings_m.columns), None)
hub_labels = []
for index, building in buildings_m.iterrows():
    name = str(building.get(hub_name_field, "Building")) if hub_name_field else "Building"
    street = str(building.get("b_addr_street", ""))
    number = str(building.get("b_addr_house_number", ""))
    address = " ".join(part for part in (street, number) if part and part != "nan").strip()
    hub_labels.append(f"{index}: {name}" + (f" — {address}" if address else ""))
if not hub_labels:
    st.error("The nPro file contains no buildings that can be selected as an energy hub.")
    st.stop()
default_hub = 0 if st.session_state.energy_hub_index is None else min(st.session_state.energy_hub_index, len(hub_labels) - 1)
selected_hub_label = hub_selector_placeholder.selectbox("Selected energy-hub building", hub_labels, index=default_hub)
selected_hub_index = int(selected_hub_label.split(":", 1)[0])
st.session_state.energy_hub_index = selected_hub_index
hub_geometry_m = buildings_m.geometry.iloc[selected_hub_index].representative_point()
hub_geometry_wgs = gpd.GeoSeries([hub_geometry_m], crs=metric_crs).to_crs(4326).iloc[0]
buildings_wgs, pipes_wgs = json_safe(buildings_m.to_crs(4326)), json_safe(pipes_m.to_crs(4326))
b = pipes_wgs.total_bounds
center = [(b[1] + b[3]) / 2, (b[0] + b[2]) / 2]

st.markdown("**1.1 Candidate placement and energy-hub selection**")
map_theme = st.selectbox("Colour pipelines by", list(THEMES))
slot = st.radio("Candidate to place", list(COLORS), horizontal=True)
interaction_mode = st.radio(
    "Map interaction", ["Place storage", "Select energy hub", "Inspect"], horizontal=True,
    help="In energy-hub mode, click a building and the nearest building becomes the selected hub.",
)
c1, c2, _ = st.columns([1, 1, 4])
if c1.button("Remove selected"):
    st.session_state.candidates.pop(slot, None)
    st.rerun()
if c2.button("Clear all"):
    st.session_state.candidates = {}
    st.rerun()

select_map = folium.Map(center, zoom_start=15, tiles=None)
add_basemaps(select_map)
add_network(select_map, buildings_wgs, pipes_wgs, map_theme, demand_col, building_demand_classes)
energy_hub_marker(select_map, hub_geometry_wgs.y, hub_geometry_wgs.x, selected_hub_label.split(":", 1)[1].strip())
current_candidates = [
    {"Candidate": name, "Latitude": item["lat"], "Longitude": item["lon"]}
    for name, item in st.session_state.candidates.items()
]
map_context_layers = nearby_map_layers(optional_layers, current_candidates, metric_crs, parcel_radius)
add_optional_layers(select_map, map_context_layers)
for name, xy in st.session_state.candidates.items():
    marker(select_map, name, xy["lat"], xy["lon"])
    folium.PolyLine(
        [[hub_geometry_wgs.y, hub_geometry_wgs.x], [xy["lat"], xy["lon"]]],
        color="#C62828", weight=4, dash_array="8 5",
        tooltip=f"Energy hub to {name}",
    ).add_to(select_map)
    if xy.get("route"):
        folium.PolyLine([[lat, lon] for lon, lat in xy["route"]], color=COLORS[name], weight=6,
                        tooltip=f"{name} manually routed connection").add_to(select_map)
Draw(export=False, position="topleft",
     draw_options={"polyline": True, "polygon": False, "rectangle": False,
                   "circle": False, "marker": False, "circlemarker": False},
     edit_options={"edit": True, "remove": True}).add_to(select_map)
folium.LayerControl(collapsed=False).add_to(select_map)
map_placeholder = st.empty()
with map_placeholder.container():
    state = st_folium(select_map, height=620, width=None, key="selection",
                      returned_objects=["last_clicked", "all_drawings"])
clicked = state.get("last_clicked") if state else None
if interaction_mode == "Place storage" and clicked:
    point = {"lat": float(clicked["lat"]), "lon": float(clicked["lng"])}
    previous = st.session_state.candidates.get(slot, {})
    point["route"] = previous.get("route")
    if previous.get("lat") != point["lat"] or previous.get("lon") != point["lon"]:
        st.session_state.candidates[slot] = point
        st.rerun()
elif interaction_mode == "Select energy hub" and clicked:
    clicked_m = gpd.GeoSeries([Point(clicked["lng"], clicked["lat"])], crs=4326).to_crs(metric_crs).iloc[0]
    distances = buildings_m.geometry.distance(clicked_m)
    nearest_hub_index = int(distances.idxmin())
    if st.session_state.energy_hub_index != nearest_hub_index:
        st.session_state.energy_hub_index = nearest_hub_index
        st.rerun()
drawings = state.get("all_drawings") if state else None
if drawings:
    lines = [item for item in drawings if item.get("geometry", {}).get("type") == "LineString"]
    if lines and slot in st.session_state.candidates:
        route = lines[-1]["geometry"]["coordinates"]
        if st.session_state.candidates[slot].get("route") != route:
            st.session_state.candidates[slot]["route"] = route
            st.rerun()
st.caption("Click to place a storage. Nearby parcels and GIS layers appear only around placed candidates. Use the polyline tool to draw a multi-bend route.")

candidate_table = pd.DataFrame([{"Candidate": n, "Latitude": v["lat"], "Longitude": v["lon"],
                                 "Manual route": bool(v.get("route"))} for n, v in st.session_state.candidates.items()])
if not candidate_table.empty:
    st.dataframe(candidate_table, hide_index=True, use_container_width=True)

analysis_requested = st.button("Analyse nPro candidates", type="primary", disabled=candidate_table.empty)

st.subheader("Step 2 · PTES geometry — dimensioned 2D and 3D design")
st.caption("The same storage volume, depth, slope, permanent perimeter and temporary working envelope are used in the drawing and GIS boundary calculation. The pump, diffuser and drainage symbols are preliminary concept-layout items.")
with st.expander("Open dimensioned 2D section, plan, 3D model and quantities", expanded=False):
    components.html(design_document, height=950, scrolling=True)
st.download_button("Download linked 2D/3D design HTML", design_document,
                   "ptes_linked_design.html", "text/html")
st.info("Preliminary design only: the layout illustrates excavation, embankment/perimeter, construction working area, pump chamber, branch pipes, diffuser zones and drainage points. Confirm final access roads, drainage, hydraulics, levels and slope stability with survey, geotechnical and specialist design inputs.")

render_thermal(design_model, power, delta_t, linked_hourly=hourly_result)

if analysis_requested:
    try:
        if not selected_scoring_criteria:
            raise ValueError("Select at least one criterion for the suitability percentage.")
        # Reuse the exact model already shown in the drawing studio above.
        flow = required_flow(power, delta_t)
        _, dn, diameter, velocity = size_connection_pipe(
            flow["volume_flow_m3_s"], min_dn=minimum_branch_dn, max_dn=maximum_branch_dn
        )
        if velocity > 2.0:
            st.warning(
                f"Required flow exceeds the screening velocity limit in DN{dn} "
                f"({velocity:.2f} m/s). Increase the maximum DN or revise power/ΔT."
            )
        elif velocity < 0.6:
            st.info(
                f"DN{dn} is enforced by the selected minimum, but its velocity is only "
                f"{velocity:.2f} m/s. Review capital cost, heat loss and controllability."
            )
        rows, map_items = [], []
        loaded_optional = sum(layer is not None for layer in optional_layers.values())
        for candidate in candidate_table.to_dict("records"):
            pt = gpd.GeoSeries([Point(candidate["Longitude"], candidate["Latitude"])], crs=4326).to_crs(metric_crs).iloc[0]
            hub_to_storage_distance = float(pt.distance(hub_geometry_m))
            excavation = translate(rotate(box(-geometry["top_length_m"]/2, -geometry["top_width_m"]/2,
                                                   geometry["top_length_m"]/2, geometry["top_width_m"]/2),
                                           rotation, origin=(0, 0)), pt.x, pt.y)
            # Square-corner envelopes make plan dimensions explicit. They are planning
            # envelopes, not a designed embankment cross-section or haul-road layout.
            embankment = excavation.buffer(embankment_width, join_style=2)
            construction = embankment.buffer(construction_clearance, join_style=2)
            parcel = containing_parcel_measurements(construction, optional_layers["Candidate parcels"])
            saved = st.session_state.candidates[candidate["Candidate"]]
            if saved.get("route"):
                route_wgs = gpd.GeoSeries([LineString(saved["route"])], crs=4326)
                route_m = route_wgs.to_crs(metric_crs).iloc[0]
                route_end = Point(route_m.coords[-1])
                nearest, final_link = nearest_pipe_connection(route_end, pipes_m)
                full_line = LineString(list(route_m.coords) + list(final_link.geometry.iloc[0].coords)[1:])
                connection = gpd.GeoDataFrame({"length_m": [full_line.length]}, geometry=[full_line], crs=metric_crs)
            else:
                nearest, connection = nearest_pipe_connection(pt, pipes_m)
            distance = float(connection["length_m"].iloc[0])
            demand = nearby_demand(pt, buildings_m, demand_col)
            demand500 = float(demand.loc[demand["Radius [m]"] == 500, "Annual demand [MWh]"].iloc[0])
            # The route length is one-way. A closed PTES connection has parallel
            # supply and return pipes, so the hydraulic circuit length is twice this value.
            pressure = darcy_weisbach_pressure_loss(2 * distance, diameter, velocity)
            def number(field):
                value = nearest.get(field)
                try:
                    return None if pd.isna(value) else float(value)
                except (TypeError, ValueError):
                    return None
            pipe_height = number("p_geo_height_m")
            candidate_height = None if pipe_height is None else pipe_height + elevation_offset
            supply_before = number("p_abs_pres_supply_sim_bar") or number("p_abs_pres_supply_bar")
            return_before = number("p_abs_pres_return_sim_bar") or number("p_abs_pres_return_bar")
            supply_at_candidate = geodetic_pressure_correction(supply_before, pipe_height, candidate_height)
            return_at_candidate = geodetic_pressure_correction(return_before, pipe_height, candidate_height)
            existing_flow_day = number("p_max_flow_rates_sim_m3_day")
            main_check = assess_main_pipe(
                None if existing_flow_day is None else existing_flow_day / 24,
                flow["volume_flow_m3_h"], number("p_diameter_DN"),
                number("p_max_power_possible_kW"), number("p_power_reserve_abs_kW"),
                power, operating_mode,
            )
            flood_overlap = overlap_area_m2(construction, optional_layers["Flood zones"])
            protected_overlap = overlap_area_m2(construction, optional_layers["Protected areas"])
            utility_crossings = intersecting_feature_count(connection.geometry.iloc[0], optional_layers["Known utilities"])
            road_distance = nearest_distance_m(construction, optional_layers["Roads"])
            groundwater = nearest_point_attribute(
                pt, optional_layers["Groundwater"],
                None if groundwater_level_field == "Not supplied" else groundwater_level_field,
            )
            parcel_fit_ratio = 1.0 if parcel["parcel_area_m2"] is None else parcel["parcel_area_m2"] / construction.area
            score, status, criteria = engineering_suitability_score(
                hub_selected=True, hub_distance_m=hub_to_storage_distance,
                land_ratio=parcel_fit_ratio, connection_m=distance,
                demand_500_mwh=demand500, reference_demand_mwh=reference_demand,
                pressure_risk=pressure["pressure_risk"], velocity_m_s=velocity,
                capacity_status=main_check["capacity_status"],
                flood_overlap_m2=flood_overlap, protected_overlap_m2=protected_overlap,
                utility_crossings=utility_crossings, loaded_optional_layers=loaded_optional,
                total_optional_layers=len(optional_layers),
                selected_criteria=selected_scoring_criteria,
            )
            criterion_scores = dict(zip(criteria["Criterion"], criteria["Criterion score [%]"]))
            criterion_weights = dict(zip(criteria["Criterion"], criteria["Applied weight [%]"]))
            row = {"Candidate": candidate["Candidate"], "Latitude": candidate["Latitude"], "Longitude": candidate["Longitude"],
                   "Score [%]": score, "Screening classification": status,
                   "Energy hub": selected_hub_label.split(":", 1)[1].strip(),
                   "Hub-to-storage distance [m]": hub_to_storage_distance,
                   "Connection [m]": distance, "Demand 500 m [MWh/year]": demand500,
                   "Storage volume [m³]": storage["volume_m3"], "Energy/cycle [MWh]": storage["energy_per_cycle_mwh"],
                   "Total pit depth [m]": depth, "Water depth [m]": design_model["h"],
                   "Freeboard [m]": freeboard,
                   "Ideal pit geometry volume [m³]": design_model["pitVolume"],
                   "Liner geometric area [m²]": design_model["linerArea"],
                   "Cover area [m²]": design_model["coverArea"],
                   "PTES excavation footprint [m²]": excavation.area,
                   "Permanent perimeter / embankment [m²]": embankment.area,
                   "Temporary construction envelope [m²]": construction.area,
                   "Temporary working area only [m²]": construction.area - embankment.area,
                   "Permanent perimeter allowance [m]": embankment_width,
                   "Temporary working clearance [m]": construction_clearance,
                   "Excavation top length [m]": geometry["top_length_m"],
                   "Excavation top width [m]": geometry["top_width_m"],
                   "Permanent envelope length [m]": geometry["top_length_m"] + 2 * embankment_width,
                   "Permanent envelope width [m]": geometry["top_width_m"] + 2 * embankment_width,
                   "Construction envelope length [m]": geometry["top_length_m"] + 2 * (embankment_width + construction_clearance),
                   "Construction envelope width [m]": geometry["top_width_m"] + 2 * (embankment_width + construction_clearance),
                   "Assessed parcel area [m²]": parcel["parcel_area_m2"],
                   "Construction outside parcel [m²]": parcel["footprint_outside_m2"],
                   "Flood-zone overlap [m²]": flood_overlap,
                   "Protected-area overlap [m²]": protected_overlap,
                   "Utility crossings [count]": utility_crossings,
                   "Nearest road [m]": road_distance,
                   "Nearest groundwater observation [m]": groundwater["distance_m"],
                   "Groundwater observed value": groundwater["value"],
                   "Top area [m²]": geometry["top_area_m2"], "Flow [m³/h]": flow["volume_flow_m3_h"],
                   "Main DN": main_check["main_dn"], "Branch DN": dn, "Velocity [m/s]": velocity,
                   "Branch sizing status": (
                       "Within screening velocity range" if 0.6 <= velocity <= 2.0
                       else "Velocity above screening limit" if velocity > 2.0
                       else "Below screening velocity; minimum DN enforced"
                   ),
                   "One-way pipe length [m]": distance, "Hydraulic circuit length [m]": 2 * distance,
                   "Round-trip pressure loss [bar]": pressure["pressure_loss_bar"], "Pressure risk": pressure["pressure_risk"],
                   "Pipe elevation [m]": pipe_height, "Candidate elevation [m]": candidate_height,
                   "Supply pressure at candidate [bar]": supply_at_candidate,
                   "Return pressure at candidate [bar]": return_at_candidate,
                   "Main flow after scenario [m³/h]": main_check["new_flow_m3_h"],
                   "Capacity status": main_check["capacity_status"],
                   **{f"Score — {name} [%]": value for name, value in criterion_scores.items()},
                   **{f"Applied weight — {name} [%]": value for name, value in criterion_weights.items()}}
            rows.append(row)
            def boundary_layer(label, shape):
                return gpd.GeoDataFrame({"Candidate": [candidate["Candidate"]], "Boundary": [label]},
                                        geometry=[shape], crs=metric_crs).to_crs(4326)
            map_items.append((row, connection.to_crs(4326),
                              boundary_layer("Excavation", excavation),
                              boundary_layer("Permanent perimeter / embankment", embankment),
                              boundary_layer("Temporary construction envelope", construction)))
    except Exception as exc:
        st.error(f"Analysis failed: {exc}")
        st.stop()

    ranking = pd.DataFrame(rows).sort_values("Score [%]", ascending=False).reset_index(drop=True)
    st.subheader("Candidate comparison")
    a, b, c = st.columns(3)
    a.metric("Best candidate", ranking.iloc[0]["Candidate"])
    b.metric("Best score", f"{ranking.iloc[0]['Score [%]']:.1f}%")
    c.metric("Shortest connection", f"{ranking['Connection [m]'].min():.1f} m")
    summary_columns = ["Candidate", "Score [%]", "Screening classification", "Energy hub",
                       "Hub-to-storage distance [m]", "Connection [m]", "Branch DN",
                       "Velocity [m/s]", "Pressure risk", "Capacity status"]
    st.dataframe(ranking[summary_columns].round(2), hide_index=True, use_container_width=True)
    st.bar_chart(ranking.set_index("Candidate")[["Score [%]"]])
    st.subheader("Stepwise engineering results")
    for _, result in ranking.iterrows():
        with st.expander(
            f"{result['Candidate']} · {result['Score [%]']:.1f}% · {result['Screening classification']}",
            expanded=result["Candidate"] == ranking.iloc[0]["Candidate"],
        ):
            st.markdown("**1. Demand and storage**")
            x1, x2, x3 = st.columns(3)
            x1.metric("Storage volume", f"{result['Storage volume [m³]']:,.0f} m³")
            x2.metric("Energy per cycle", f"{result['Energy/cycle [MWh]']:,.1f} MWh")
            x3.metric("Demand within 500 m", f"{result['Demand 500 m [MWh/year]']:,.0f} MWh/year")
            st.markdown("**2. Energy hub and network route**")
            st.write(f"Energy hub: {result['Energy hub']}")
            st.write(f"Hub → storage: {result['Hub-to-storage distance [m]']:.1f} m")
            st.write(f"Storage → existing network: {result['Connection [m]']:.1f} m")
            st.markdown("**3. Hydraulic compatibility**")
            h1, h2, h3, h4 = st.columns(4)
            h1.metric("Branch", f"DN {int(result['Branch DN'])}")
            h2.metric("Flow", f"{result['Flow [m³/h]']:.1f} m³/h")
            h3.metric("Velocity", f"{result['Velocity [m/s]']:.2f} m/s")
            h4.metric("Circuit loss", f"{result['Round-trip pressure loss [bar]']:.2f} bar")
            st.write(f"{result['Branch sizing status']} · {result['Capacity status']}")
            st.markdown("**4. Land, construction and constraints**")
            land_table = pd.DataFrame({
                "Result": ["Excavation rim footprint", "Permanent perimeter / embankment", "Temporary construction envelope",
                           "Permanent envelope dimensions", "Construction-envelope dimensions",
                           "Outside parcel", "Flood overlap", "Protected-area overlap",
                           "Utility crossings", "Nearest groundwater observation"],
                "Value": [f"{result['PTES excavation footprint [m²]']:.0f} m²",
                          f"{result['Permanent perimeter / embankment [m²]']:.0f} m² ({result['Permanent perimeter allowance [m]']:.1f} m allowance)",
                          f"{result['Temporary construction envelope [m²]']:.0f} m² ({result['Temporary working clearance [m]']:.1f} m working clearance)",
                          f"{result['Permanent envelope length [m]']:.1f} × {result['Permanent envelope width [m]']:.1f} m",
                          f"{result['Construction envelope length [m]']:.1f} × {result['Construction envelope width [m]']:.1f} m",
                          "Not assessed" if pd.isna(result['Construction outside parcel [m²]']) else f"{result['Construction outside parcel [m²]']:.0f} m²",
                          "Not assessed" if pd.isna(result['Flood-zone overlap [m²]']) else f"{result['Flood-zone overlap [m²]']:.0f} m²",
                          "Not assessed" if pd.isna(result['Protected-area overlap [m²]']) else f"{result['Protected-area overlap [m²]']:.0f} m²",
                          "Not assessed" if pd.isna(result['Utility crossings [count]']) else str(int(result['Utility crossings [count]'])),
                          "Not assessed" if pd.isna(result['Nearest groundwater observation [m]']) else f"{result['Nearest groundwater observation [m]']:.0f} m"],
            })
            st.dataframe(land_table, hide_index=True, use_container_width=True)
            st.markdown("**5. Weighted suitability criteria**")
            score_columns = [column for column in ranking.columns if column.startswith("Score —")]
            score_table = pd.DataFrame({
                "Criterion": [column.removeprefix("Score — ").removesuffix(" [%]") for column in score_columns],
                "Score [%]": [result[column] for column in score_columns],
                "Applied weight [%]": [
                    result[f"Applied weight — {column.removeprefix('Score — ').removesuffix(' [%]')} [%]"]
                    for column in score_columns
                ],
            })
            st.dataframe(score_table.round(1), hide_index=True, use_container_width=True)
    result_map = folium.Map(center, zoom_start=15, tiles=None)
    add_basemaps(result_map)
    add_network(result_map, buildings_wgs, pipes_wgs, map_theme, demand_col, building_demand_classes)
    energy_hub_marker(result_map, hub_geometry_wgs.y, hub_geometry_wgs.x,
                      selected_hub_label.split(":", 1)[1].strip())
    add_optional_layers(result_map, map_context_layers)
    for result, connection, excavation, embankment, construction in map_items:
        name, color = result["Candidate"], COLORS[result["Candidate"]]
        folium.GeoJson(construction, name=f"{name} temporary construction envelope",
                       style_function=lambda _: {"color": "#555", "weight": 2, "dashArray": "7 5",
                                                 "fillColor": "#999", "fillOpacity": .10}).add_to(result_map)
        folium.GeoJson(embankment, name=f"{name} permanent perimeter / embankment",
                       style_function=lambda _: {"color": "#B7791F", "weight": 2,
                                                 "fillColor": "#D69E2E", "fillOpacity": .15}).add_to(result_map)
        folium.GeoJson(connection, name=f"{name} connection", style_function=lambda _, col=color: {"color": col, "weight": 6},
                       tooltip=f"{name}: {result['Connection [m]']:.1f} m").add_to(result_map)
        folium.PolyLine(
            [[hub_geometry_wgs.y, hub_geometry_wgs.x], [result["Latitude"], result["Longitude"]]],
            color="#C62828", weight=4, dash_array="8 5",
            tooltip=f"Energy hub to {name}: {result['Hub-to-storage distance [m]']:.1f} m",
        ).add_to(result_map)
        folium.GeoJson(excavation, name=f"{name} excavation boundary", style_function=lambda _, col=color: {"color": col, "weight": 3, "fillColor": col, "fillOpacity": .38}).add_to(result_map)
        marker(result_map, name, result["Latitude"], result["Longitude"], result)
    best = ranking.iloc[0]
    summary_panel = f"""<div style='position:fixed;bottom:24px;right:24px;z-index:9999;background:white;
    padding:12px 14px;border:1px solid #555;border-left:5px solid #d97706;border-radius:5px;
    max-width:310px;font-size:12px;box-shadow:0 2px 8px rgba(0,0,0,.18)'>
    <b style='font-size:14px'>♨ PTES Engineering Screening</b><br>
    <span style='color:#555'>Kuruba Pujari Gopal<br>Rother und Partner Ingenieurgesellschaft</span><hr style='margin:7px 0'>
    Leading candidate: {best['Candidate']}<br>
    Storage volume: {best['Storage volume [m³]']:.0f} m³<br>
    Excavation footprint: {best['PTES excavation footprint [m²]']:.0f} m²<br>
    Permanent perimeter / embankment: {best['Permanent perimeter / embankment [m²]']:.0f} m²<br>
    Temporary construction envelope: {best['Temporary construction envelope [m²]']:.0f} m²<br>
    Connection: {best['Connection [m]']:.1f} m · branch DN {best['Branch DN']}<br>
    Score criteria used: {len(selected_scoring_criteria)} of {len(SCORING_CRITERIA)}<br>
    <span style='font-size:10px'>{' · '.join(selected_scoring_criteria)}</span><br>
    Optional authority layers loaded: {loaded_optional}/{len(optional_layers)}<br>
    Missing layers remain not assessed.</div>"""
    result_map.get_root().html.add_child(folium.Element(summary_panel))
    folium.LayerControl(collapsed=False).add_to(result_map)
    # streamlit-folium adds its own synchronisation map object to the Folium root.
    # Capture the standalone export first so the downloaded file contains one map only.
    standalone_html = result_map.get_root().render().encode("utf-8")
    map_placeholder.empty()
    with map_placeholder.container():
        st_folium(result_map, height=700, width=None, key="results", returned_objects=[])
    st.download_button("Download interactive HTML map", standalone_html, "ptes_candidate_comparison_map.html", "text/html")
    st.download_button("Download comparison CSV", ranking.to_csv(index=False).encode("utf-8"), "ptes_candidate_comparison.csv", "text/csv")
    st.download_button("Download GIS data register", register.to_csv(index=False).encode("utf-8"), "ptes_data_register.csv", "text/csv")
    active_inputs = pd.DataFrame([
        {"Input": "PTES water volume [m³]", "Active value": storage["volume_m3"]},
        {"Input": "Hot reference temperature [°C]", "Active value": tmax},
        {"Input": "Cold reference temperature [°C]", "Active value": tmin},
        {"Input": "Solar thermal capacity [kWth]", "Active value": solar_thermal_kw},
        {"Input": "Solar net specific yield [kWhth/kWth/year]", "Active value": solar_specific_yield},
        {"Input": "BHKW 1 electrical / thermal [kW]", "Active value": f"{bhkw1_electrical_kw} / {bhkw1_thermal_kw}"},
        {"Input": "BHKW 2 electrical / thermal [kW]", "Active value": f"{bhkw2_electrical_kw} / {bhkw2_thermal_kw}"},
        {"Input": "Summer BHKW price threshold [€/MWh]", "Active value": price_threshold},
        {"Input": "BHKW-to-PTES charging months", "Active value": ", ".join(bhkw_charge_labels)},
        {"Input": "BHKW direct-network months", "Active value": ", ".join(bhkw_direct_labels)},
        {"Input": "PTES discharge months", "Active value": ", ".join(ptes_discharge_labels)},
        {"Input": "PTES cycle view", "Active value": storage_cycle_view},
        {"Input": "Initial PTES state of charge [%]", "Active value": hourly_initial_soc * 100},
    ])
    st.download_button("Download complete PTES project Excel workbook",
                       project_excel_workbook(active_inputs, energy_frame, hourly_result, design_model,
                                              ranking, register, weather_seasonal, weather_monthly),
                       "ptes_complete_project_workbook.xlsx",
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

st.divider()
st.caption("Preliminary screening only. Validate topology, pumps, boundary pressures, pipe roughness, fittings and operating cases in nPro.")
